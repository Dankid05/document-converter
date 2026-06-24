import os
from django.shortcuts import render
from django.http import FileResponse
from django.conf import settings

def index(request):
    return render(request, 'index.html')

def convert(request):
    if request.method == 'POST':
        file = request.FILES.get('file')
        conversion_type = request.POST.get('conversion_type')
        typed_text = request.POST.get('typed_text', '').strip()

        # Validate input
        if not file and not typed_text:
            return render(request, 'index.html', {'error': 'Please upload a file or type some text!'})

        # Save uploaded file if exists
        upload_path = None
        if file:
            upload_dir = os.path.join(settings.MEDIA_ROOT, 'uploads')
            os.makedirs(upload_dir, exist_ok=True)
            upload_path = os.path.join(upload_dir, file.name)

            with open(upload_path, 'wb+') as f:
                for chunk in file.chunks():
                    f.write(chunk)

        # Output folder
        output_dir = os.path.join(settings.MEDIA_ROOT, 'converted')
        os.makedirs(output_dir, exist_ok=True)

        output_path = None

        try:
            # Word to PDF
          if conversion_type == 'word_to_pdf':
            import subprocess
            import shutil
            output_path = os.path.join(output_dir, file.name.replace('.docx', '.pdf'))
            
            # Find libreoffice wherever it is
            libreoffice_path = shutil.which('libreoffice') or shutil.which('soffice')
            
            if not libreoffice_path:
                raise Exception("LibreOffice is not installed on this server")
            
            result = subprocess.run([
                libreoffice_path, '--headless', '--convert-to', 'pdf',
                '--outdir', output_dir, upload_path
            ], capture_output=True, text=True)
            
            if result.returncode != 0:
                raise Exception(f"LibreOffice error: {result.stderr}")

            # Text to Word
            elif conversion_type == 'text_to_word':
                from docx import Document
                output_filename = file.name.replace('.txt', '.docx') if file else 'converted.docx'
                output_path = os.path.join(output_dir, output_filename)
                doc = Document()
                

                if typed_text:
                    for line in typed_text.splitlines():
                        doc.add_paragraph(line)
                elif file:
                    with open(upload_path, 'r') as f:
                        for line in f:
                            doc.add_paragraph(line.strip())

                doc.save(output_path)

            # PDF to Word
            elif conversion_type == 'pdf_to_word':
                from pdf2docx import Converter
                output_path = os.path.join(output_dir, file.name.replace('.pdf', '.docx'))
                cv = Converter(upload_path)
                cv.convert(output_path)
                cv.close()

            if output_path and os.path.exists(output_path):
                return FileResponse(
                    open(output_path, 'rb'),
                    as_attachment=True,
                    filename=os.path.basename(output_path)
                )

        except Exception as e:
            return render(request, 'index.html', {'error': f'Conversion failed: {str(e)}'})

    return render(request, 'index.html')